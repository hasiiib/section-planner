"""Export identical day-grouped timetables in all supported formats."""
import csv
from pathlib import Path
from .model import DAYS
from .scheduler import verify,time_label

HEADERS=['Day','Class Period','Course Code','Course Title','Section','Faculty','Room']

def rows(routine):
    problems=verify(routine)
    if problems:raise ValueError(f'Invalid routine cannot be exported: {problems[:2]}')
    entries=sorted(((DAYS.index(m.day),m.start,m.day,f'{time_label(m.start)} – {time_label(m.end)}',s.code,s.title,s.section,s.faculty,m.room) for s in routine for m in s.meetings))
    return [list(e[2:]) for e in entries]

def spans(data):
    i=0
    while i<len(data):
        j=i+1
        while j<len(data) and data[j][0]==data[i][0]:j+=1
        yield i,j-i
        i=j

def export_csv(routine,path):
    data=rows(routine)
    with open(path,'w',newline='',encoding='utf-8-sig') as f:
        out=csv.writer(f);out.writerow(['Class Routine']);out.writerow(HEADERS)
        for start,length in spans(data):
            for i in range(start,start+length):out.writerow(data[i] if i==start else ['']+data[i][1:])

def export_xlsx(routine,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    data=rows(routine);wb=Workbook();ws=wb.active;ws.title='Class Routine'
    ws.merge_cells('A1:G1');ws['A1']='Class Routine';ws['A1'].font=Font(size=20,bold=True,color='F47721');ws['A1'].alignment=Alignment(horizontal='center')
    ws.append(HEADERS)
    for item in data:ws.append(item)
    for c in ws[2]:c.fill=PatternFill('solid',fgColor='F47721');c.font=Font(bold=True,color='FFFFFF');c.alignment=Alignment(horizontal='center')
    for start,length in spans(data):
        if length>1:ws.merge_cells(start_row=start+3,start_column=1,end_row=start+length+2,end_column=1)
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment=Alignment(vertical='center',wrap_text=True)
            c.border=Border(bottom=Side(style='hair',color='E6E9EF'))
    for i,w in enumerate([13,25,19,52,13,34,27],1):ws.column_dimensions[get_column_letter(i)].width=w
    for i in range(3,3+len(data)):ws.row_dimensions[i].height=40
    ws.freeze_panes='A3';wb.save(path)

def export_pdf(routine,path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3,landscape
    from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer,KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    data=rows(routine)
    doc=SimpleDocTemplate(str(path),pagesize=landscape(A3),leftMargin=25,rightMargin=25,topMargin=25,bottomMargin=25)
    style=ParagraphStyle('cell',fontName='Helvetica',fontSize=8,leading=11)
    body=[[Paragraph(str(v),style) for v in HEADERS]]+[[Paragraph(str(v),style) for v in record] for record in data]
    grid=Table(body,colWidths=[57,139,95,235,65,182,104],repeatRows=1,hAlign='CENTER')
    commands=[('BACKGROUND',(0,0),(-1,0),colors.HexColor('#F47721')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('BOTTOMPADDING',(0,0),(-1,-1),11),('TOPPADDING',(0,0),(-1,-1),11),('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor('#D65F16')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#FFF7F0')])]
    for start,length in spans(data):
        if length>1:commands.append(('SPAN',(0,start+1),(0,start+length)))
        commands.append(('LINEBELOW',(0,start+length),(-1,start+length),.45,colors.HexColor('#DFE4EC')))
        for row in range(start+1,start+length):commands.append(('LINEBELOW',(1,row),(-1,row),.3,colors.HexColor('#E6E9EF')))
    grid.setStyle(TableStyle(commands))
    styles=getSampleStyleSheet();doc.build([Paragraph('Class Routine',styles['Title']),Spacer(1,14),grid])

def export_png(routine,path):
    from PIL import Image,ImageDraw,ImageFont
    data=rows(routine);widths=[100,235,165,485,105,320,230];w=sum(widths)+40
    font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    try:font=ImageFont.truetype(font_path,18);bold=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',18);titlefont=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',32)
    except OSError:font=bold=titlefont=ImageFont.load_default()
    def wrap(text,limit,font):
        text=str(text);words=text.split();lines=['']
        for word in words:
            trial=(lines[-1]+' '+word).strip()
            if ImageDraw.Draw(Image.new('RGB',(1,1))).textbbox((0,0),trial,font=font)[2]>limit and lines[-1]:lines.append(word)
            else:lines[-1]=trial
        return lines
    heights=[max(46,17+27*max(len(wrap(v,widths[j]-16,font)) for j,v in enumerate(row))) for row in data]
    h=130+sum(heights)+12;im=Image.new('RGB',(w,h),'#FFFFFF');d=ImageDraw.Draw(im)
    d.text((24,15),'Class Routine',font=titlefont,fill='#222B39')
    d.rounded_rectangle((14,78,w-14,120),radius=9,fill='#F47721')
    x=20
    for j,hd in enumerate(HEADERS):d.text((x+8,88),hd,font=bold,fill='white');x+=widths[j]
    y=120
    for i,(row,height) in enumerate(zip(data,heights)):
        d.rectangle((14,y,w-14,y+height),fill='#FFFFFF' if i%2 else '#FFF8F3')
        x=20
        for j,value in enumerate(row):
            if j:
                for k,line in enumerate(wrap(value,widths[j]-16,font)):d.text((x+8,y+9+26*k),line,font=font,fill='#263043')
            x+=widths[j]
        d.line((120,y+height,w-14,y+height),fill='#E5E9F0',width=1)
        y+=height
    # Vertically center each merged day label and suppress inner borders in day column.
    y0=120
    for start,count in spans(data):
        start_y=y0+sum(heights[:start]);total=sum(heights[start:start+count]);day=data[start][0]
        d.text((30,start_y+total/2-10),day,font=bold,fill='#263043')
        d.line((14,start_y+total,w-14,start_y+total),fill='#DEE4EC',width=2)
    im.save(path)

def export(routine,path):
    methods={'.pdf':export_pdf,'.png':export_png,'.xlsx':export_xlsx,'.csv':export_csv}
    ext=Path(path).suffix.lower()
    if ext not in methods:raise ValueError('Unsupported export format: '+ext)
    return methods[ext](routine,path)
