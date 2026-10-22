"""Strict structured-file import. Missing schedule facts are never invented."""
import csv
import re
from pathlib import Path
from .parser import parse_pdf, detect_trimester, time_range
from .model import Section, Meeting, DAYS

HEADERS = ('program','course code','title','section','room1','room2','day1','day2','time1','time2','faculty name','faculty initial','credit')
ALIASES = {'course':'course code','coursecode':'course code','course title':'title','faculty':'faculty name','facultyinitial':'faculty initial','facultyname':'faculty name','room':'room1','day':'day1','time':'time1','sl':'serial','serial number':'serial'}
SUPPORTED = ('.pdf','.csv','.xlsx','.xls','.docx','.doc')

def normalized(value):
    name = re.sub(r'\s+',' ',str(value or '').strip().lower().replace('_',' '))
    return ALIASES.get(name,name)

def _table_to_sections(table):
    if not table: raise ValueError('This file has no rows.')
    headers=[normalized(x) for x in table[0]]
    needed={'program','course code','title','section','day1','time1'}
    if not needed.issubset(headers):
        raise ValueError('Not a valid course-offering table. Required columns: '+', '.join(sorted(needed- set(headers))))
    result=[];warnings=[];seen=set()
    for idx,values in enumerate(table[1:],2):
        if not any(str(x or '').strip() for x in values):continue
        d={key:str(values[pos] if pos<len(values) and values[pos] is not None else '').strip() for pos,key in enumerate(headers)}
        program=d.get('program','').upper()
        code=' '.join(d.get('course code','').upper().split())
        title=d.get('title','');section=d.get('section','').strip()
        conditional=bool(re.search(r'\bif\s+required\b',section,re.I))
        section=re.sub(r'\s*\(?if\s+required\)?','',section,flags=re.I).strip()
        if not (program and re.fullmatch(r'[A-Z]+\s+\d+[A-Z]?',code) and title and section):
            warnings.append(f'Row {idx}: missing/invalid program, code, title or section (skipped)');continue
        meetings=[];problems=[]
        for suffix in ('1','2'):
            day=d.get('day'+suffix,'').title()[:3]
            timing=d.get('time'+suffix,'')
            if not day and not timing:continue
            if day not in DAYS or not timing:
                problems.append(f'Day{suffix}/Time{suffix} incomplete');continue
            try:
                start,end=time_range(timing)
                meetings.append(Meeting(day,start,end,d.get('room'+suffix,'')))
            except ValueError:
                problems.append(f'Invalid Time{suffix}: {timing}')
        if not meetings:
            warnings.append(f'Row {idx}: no complete day/time slot (skipped)');continue
        key=(program,code,section)
        if key in seen:
            warnings.append(f'Row {idx}: duplicate section {key} (skipped)');continue
        seen.add(key)
        result.append(Section(program,code,title,section,d.get('faculty name','') or 'TBA',d.get('faculty initial',''),d.get('credit',''),tuple(meetings),idx,(d.get('notes','')+' (If Required)' if conditional else d.get('notes','')),tuple(problems)))
        if problems:warnings.append(f'Row {idx}: '+', '.join(problems))
    if not result:raise ValueError('No valid course sections with complete class periods were found.')
    return result,warnings,1

def _read_csv(path):
    with open(path,encoding='utf-8-sig',newline='') as stream:
        sample=stream.read(4096);stream.seek(0)
        try:dialect=csv.Sniffer().sniff(sample)
        except csv.Error:dialect=csv.excel
        return list(csv.reader(stream,dialect))

def _read_xlsx(path):
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        for ws in wb:
            data=list(ws.values)
            for i,row in enumerate(data[:30]):
                if {'program','course code','section'}.issubset({normalized(x) for x in row}):return data[i:]
    finally:wb.close()
    raise ValueError('No worksheet contains the required course-offering header.')

def _read_xls(path):
    import xlrd
    book=xlrd.open_workbook(path)
    for sheet in book.sheets():
        data=[sheet.row_values(i) for i in range(sheet.nrows)]
        for i,row in enumerate(data[:30]):
            if {'program','course code','section'}.issubset({normalized(x) for x in row}):return data[i:]
    raise ValueError('No worksheet contains the required course-offering header.')

def _read_docx(path):
    from docx import Document
    doc=Document(path)
    for table in doc.tables:
        data=[[cell.text.strip() for cell in row.cells] for row in table.rows]
        for i,row in enumerate(data[:30]):
            if {'program','course code','section'}.issubset({normalized(x) for x in row}):return data[i:]
    raise ValueError('No Word table contains the required course-offering header.')

def _read_doc(path):
    # Legacy binary .doc needs LibreOffice; no admin install attempted.
    import shutil,subprocess,tempfile
    lo=shutil.which('libreoffice') or shutil.which('soffice')
    if not lo:raise ValueError('Old .doc files require LibreOffice installed locally. Save as .docx or install LibreOffice yourself.')
    with tempfile.TemporaryDirectory() as directory:
        process=subprocess.run([lo,'-env:UserInstallation=file://'+directory+'/profile','--headless','--convert-to','docx','--outdir',directory,str(path)],capture_output=True,text=True,timeout=90)
        converted=Path(directory)/(Path(path).stem+'.docx')
        if process.returncode or not converted.exists():raise ValueError('Could not convert the .doc file to .docx. Please save it as .docx.')
        return _read_docx(converted)

def import_offerings(path):
    path=Path(path)
    ext=path.suffix.lower()
    if ext not in SUPPORTED:raise ValueError('Unsupported file type. Supported: PDF, DOC, DOCX, XLS, XLSX and CSV.')
    if ext=='.pdf':
        records,warnings,pages=parse_pdf(path)
        if not records:raise ValueError('No valid course offerings found in this PDF.')
        return records,warnings,pages,detect_trimester(path)
    readers={'.csv':_read_csv,'.xlsx':_read_xlsx,'.xls':_read_xls,'.docx':_read_docx,'.doc':_read_doc}
    records,warnings,pages=_table_to_sections(readers[ext](path))
    return records,warnings,pages,None
