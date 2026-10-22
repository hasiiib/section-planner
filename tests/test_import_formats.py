import csv
from pathlib import Path
import pytest
from uiu.importers import import_offerings, SUPPORTED
from uiu.exports import export, rows, spans
from uiu.model import Section,Meeting

HEADER=['Program','Course Code','Title','Section','Day1','Time1','Room1','Day2','Time2','Room2','Faculty Name','Credit']
DATA=[['BSCSE','CSE 1111','Structured Programming Language','A','Sat','08:30:AM - 09:50:AM','301','Tue','08:30:AM - 09:50:AM','302','TBA','3'],['BSDS','DS 1201','Data Science','B','Sun','11:11:AM - 12:30:PM','400','','','','TBA','3']]

def check_import(path):
    sections,warnings,_,_=import_offerings(path)
    assert len(sections)==2
    assert {s.program for s in sections}=={'BSCSE','BSDS'}
    assert len(sections[0].meetings)==2
    assert len(sections[1].meetings)==1
    assert sections[1].faculty=='TBA'
    return sections

def test_csv_import(tmp_path):
    path=tmp_path/'offering.csv'
    with path.open('w',newline='') as f:csv.writer(f).writerows([HEADER]+DATA)
    check_import(path)

def test_excel_import(tmp_path):
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    for row in [HEADER]+DATA:ws.append(row)
    path=tmp_path/'offering.xlsx';wb.save(path)
    check_import(path)

def test_word_import(tmp_path):
    from docx import Document
    d=Document();t=d.add_table(rows=1,cols=len(HEADER))
    for c,value in zip(t.rows[0].cells,HEADER):c.text=value
    for values in DATA:
        for c,value in zip(t.add_row().cells,values):c.text=value
    path=tmp_path/'offering.docx';d.save(path)
    check_import(path)

def test_unsupported_import(tmp_path):
    path=tmp_path/'notes.txt';path.write_text('anything')
    with pytest.raises(ValueError,match='Unsupported'):import_offerings(path)

def test_invalid_csv(tmp_path):
    path=tmp_path/'wrong.csv';path.write_text('a,b,c\n1,2,3\n')
    with pytest.raises(ValueError,match='Not a valid'):import_offerings(path)

def test_no_schedule_csv(tmp_path):
    path=tmp_path/'no_days.csv'
    with path.open('w',newline='') as f:csv.writer(f).writerows([HEADER,[*DATA[0][:4],'','','','','','','TBA','3']])
    with pytest.raises(ValueError,match='No valid'):import_offerings(path)

def test_merged_day_export(tmp_path):
    s1=Section('BSCSE','CSE 1111','Programming','A','TBA','','3',(Meeting('Sat',510,590,'301'),Meeting('Tue',510,590,'301')))
    s2=Section('BSCSE','CSE 1112','Programming Lab','A','TBA','','1',(Meeting('Sat',600,700,'302'),))
    path=tmp_path/'routine.xlsx';export([s1,s2],path)
    from openpyxl import load_workbook
    ws=load_workbook(path).active
    assert ws['A1'].value=='Class Routine'
    assert 'A3:A4' in [str(r) for r in ws.merged_cells.ranges]
    assert len(rows([s1,s2]))==3
    assert list(spans(rows([s1,s2])))==[(0,2),(2,1)]
